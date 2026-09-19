"""
Layer 2 - review engine (screening heuristics, not legal conclusions).

Input : normalized_transactions.csv (layer 1) + case_config.xlsx
Output: review_output.xlsx with sheets Transactions / Review_Items / Counterparty_Summary / Case_Summary

Each rule is a function (df, ctx) -> list[Hit]. Hits are accumulated per transaction:
several rules may flag the same transaction; Rule_IDs and Review_Reasons are joined with ';'.

Usage: python review_engine.py --transactions ../output/normalized_transactions.csv \
                               --config ../synthetic/data/case_config.xlsx --out ../output/review_output.xlsx
"""
from __future__ import annotations
import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path
import pandas as pd

# ----------------------------------------------------------------------------- context
@dataclass
class Thresholds:
    high_value: float = 10_000
    cash_withdrawal: float = 1_000
    cash_repeat_count: int = 5
    cash_repeat_days: int = 30
    round_step: float = 1_000
    round_min: float = 5_000
    new_cp_first_payment: float = 5_000
    # Analytical recency window only; it is not a statutory look-back period.
    # A counterparty is treated as "new" by R06 only when its first observed
    # outgoing payment falls within this many days before insolvency.
    # 0 disables the recency restriction and uses the whole analysis period.
    new_cp_window_days: int = 180

@dataclass
class Context:
    company: str
    appointment_date: pd.Timestamp
    insolvency_date: pd.Timestamp
    analysis_start: pd.Timestamp        # computed from anchor + lookback, or override
    currency: str
    known_parties: pd.DataFrame          # Name, Relationship, Sort code, Account no., Notes, Name_Norm
    thresholds: Thresholds = field(default_factory=Thresholds)
    procedure: str = ""
    petition_date: pd.Timestamp | None = None
    resolution_date: pd.Timestamp | None = None
    analysis_anchor: str = "Appointment"          # Appointment / Petition / Resolution / Insolvency / Custom
    anchor_date: pd.Timestamp | None = None       # the date the lookback is measured from
    lookback_months: int | None = None
    review_output_mode: str = "All flagged"       # "All flagged" | "Analysis period only"
    start_overridden: bool = False

ANCHOR_FIELD = {"appointment": "appointment_date", "petition": "petition_date",
                "resolution": "resolution_date", "insolvency": "insolvency_date"}

def _ts(v):
    return None if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "NaT") else pd.Timestamp(v)

def resolve_analysis_period(case: pd.Series) -> tuple[str, pd.Timestamp, int | None, pd.Timestamp, bool]:
    """Analyst-defined period: anchor + lookback months, or an explicit start override.
    Returns (anchor_name, anchor_date, lookback_months, start, start_overridden).
    The analyst decides the scope; the tool computes the boundary. Not a statutory look-back.
    - Analysis_Anchor: Appointment / Petition / Resolution / Insolvency / Custom (Custom uses Analysis_Anchor_Date)
    - Analysis_Period_Start_Override changes ONLY the start; it never changes the anchor."""
    anchor_name = str(case.get("Analysis_Anchor", "Appointment") or "Appointment").strip()
    override = _ts(case.get("Analysis_Period_Start_Override")) or _ts(case.get("Analysis_Period_Start"))
    months_raw = case.get("Analysis_Lookback_Months")
    months = int(float(months_raw)) if months_raw is not None and str(months_raw).strip() not in ("", "nan") else None
    dates = {"appointment": _ts(case.get("Appointment date")), "petition": _ts(case.get("Petition_Date")),
             "resolution": _ts(case.get("Resolution_Date")), "insolvency": _ts(case.get("Insolvency date")),
             "custom": _ts(case.get("Analysis_Anchor_Date"))}
    key = anchor_name.lower()
    if key not in dates:
        raise ValueError(f"Analysis_Anchor='{anchor_name}' is not one of Appointment/Petition/Resolution/Insolvency/Custom")
    anchor_date = dates[key]
    if anchor_date is None:
        need = "Analysis_Anchor_Date" if key == "custom" else f"the {anchor_name} date"
        raise ValueError(f"Analysis_Anchor='{anchor_name}' but {need} is not set in case_config")
    if months is not None and months <= 0:
        raise ValueError("Analysis_Lookback_Months must be a positive number of months")
    if override is None and months is None:
        raise ValueError("Set Analysis_Lookback_Months or Analysis_Period_Start_Override")
    start = override if override is not None else anchor_date - pd.DateOffset(months=months)
    if start >= anchor_date:
        raise ValueError(f"Analysis period start {start:%d %b %Y} must be before the anchor date {anchor_date:%d %b %Y}")
    return anchor_name.capitalize(), anchor_date, months, start, override is not None

def _norm(s) -> str:
    s = "" if pd.isna(s) else str(s).upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\b(LTD|LIMITED|PLC|LLP|LLC|CO|COMPANY)\b", "", s)
    return re.sub(r"\s+", " ", s).strip()

def load_context(config_path: Path) -> Context:
    case = pd.read_excel(config_path, sheet_name="Case").set_index("Field")["Value"]
    kp = pd.read_excel(config_path, sheet_name="Known_Parties", dtype=str).fillna("")
    kp["Name_Norm"] = kp["Name"].map(_norm)
    th = Thresholds()
    try:
        t = pd.read_excel(config_path, sheet_name="Thresholds").set_index("Threshold")["Value"]
        th.high_value = float(t.get("High value", th.high_value))
        th.cash_withdrawal = float(t.get("Cash withdrawal", th.cash_withdrawal))
        th.round_step = float(t.get("Round amount step", th.round_step))
        th.round_min = float(t.get("Round amount minimum", th.round_min))
        th.new_cp_first_payment = float(t.get("New counterparty first payment", th.new_cp_first_payment))
        th.new_cp_window_days = int(float(t.get("New counterparty window (days)", th.new_cp_window_days)))
    except Exception:
        pass
    anchor_name, anchor_date, months, start, overridden = resolve_analysis_period(case)
    return Context(
        company=str(case["Company name"]),
        appointment_date=pd.Timestamp(case["Appointment date"]),
        insolvency_date=pd.Timestamp(case["Insolvency date"]),
        analysis_start=pd.Timestamp(start),
        currency=str(case.get("Currency", "GBP")),
        known_parties=kp, thresholds=th,
        procedure=str(case.get("Procedure", "") or ""),
        petition_date=_ts(case.get("Petition_Date")), resolution_date=_ts(case.get("Resolution_Date")),
        analysis_anchor=anchor_name, anchor_date=anchor_date, lookback_months=months, start_overridden=overridden,
        review_output_mode=str(case.get("Review_Output_Mode", "All flagged") or "All flagged"),
    )

def load_transactions(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype={"Account": str, "Counterparty_Account": str, "Counterparty_Sort_Code": str})
    df["Date"] = pd.to_datetime(df["Date"])
    for c in ("Money_In", "Money_Out", "Amount_Signed", "Balance"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0).abs() if c != "Amount_Signed" else pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    df["Money_Out"] = df["Money_Out"].abs()
    df["Counterparty_Norm"] = df["Counterparty_Norm"].fillna("").astype(str)
    df["Is_Own_Account_Transfer"] = df["Is_Own_Account_Transfer"].astype(str).str.lower().eq("true")
    return df

# ----------------------------------------------------------------------------- rules
@dataclass
class Hit:
    txn_id: str
    rule_id: str
    reason: str

def fmt_money(x: float, cur: str = "£") -> str:
    return f"{cur}{x:,.0f}" if float(x).is_integer() else f"{cur}{x:,.2f}"

def fmt_total(x: float, cur: str = "£") -> str:
    return f"{cur}{x:,.2f}"

def currency_symbol(code: str) -> str:
    return {"GBP": "£", "EUR": "€", "USD": "$"}.get(str(code).upper(), f"{code} ")

def _known_party_maps(ctx: Context):
    """Return exact-name maps for known parties.

    Matching is deliberately exact after conservative normalisation.  We do
    not resolve initials (for example ``J SMITH``) to a person when more than
    one known party could plausibly match.
    """
    kp = ctx.known_parties.copy()
    kp["Name_Norm"] = kp["Name_Norm"].map(_norm)
    kp = kp[kp["Name_Norm"] != ""].drop_duplicates("Name_Norm", keep="first")
    by_name = kp.set_index("Name_Norm")
    return by_name

def rule_R01_connected_party(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Outgoing payment to an exactly matched known connected party.

    Own accounts, HMRC and the secured lender are known parties for context,
    but are not treated as connected-party review hits by this rule.
    """
    kp = _known_party_maps(ctx)
    excluded_relationships = {"OWN ACCOUNT", "HMRC", "SECURED LENDER"}
    eligible = kp[~kp["Relationship"].str.upper().isin(excluded_relationships)]
    if eligible.empty:
        return []

    norm = df["Counterparty_Norm"].map(_norm)
    m = (df["Money_Out"] > 0) & (~df["Is_Own_Account_Transfer"]) & norm.isin(eligible.index)
    hits = []
    for idx, r in df[m].iterrows():
        n = _norm(r["Counterparty_Norm"])
        k = eligible.loc[n]
        hits.append(Hit(r["Txn_ID"], "R01",
                        f"Payment to known connected party: {k['Name']} ({k['Relationship']})"))
    return hits

def rule_R03_high_value(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """High-value transaction above threshold (excludes own-account transfers and known
    routine payees: HMRC, secured lender). Applies to money out only."""
    th = ctx.thresholds.high_value
    routine = set(ctx.known_parties.loc[ctx.known_parties["Relationship"].isin(["HMRC", "Secured lender", "Own account"]), "Name_Norm"])
    m = (df["Money_Out"] >= th) & (~df["Is_Own_Account_Transfer"]) & (~df["Counterparty_Norm"].map(_norm).isin(routine))
    sym = currency_symbol(ctx.currency)
    return [Hit(r.Txn_ID, "R03", f"High-value transaction above {fmt_money(th, sym)}") for r in df[m].itertuples()]

def rule_R10_description_terms(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Outgoing payment whose description explicitly indicates loan/dividend/director."""
    text = df["Description"].fillna("").astype(str)
    # Whole-word matching avoids accidental hits inside unrelated words.
    m = (df["Money_Out"] > 0) & (~df["Is_Own_Account_Transfer"]) & \
        text.str.contains(r"\b(?:loan|dividend|director)\b", case=False, regex=True)
    return [Hit(r.Txn_ID, "R10", "Description indicates loan/dividend/director payment")
            for r in df[m].itertuples()]

def _cash_mask(df: pd.DataFrame) -> pd.Series:
    typ = df["Type"].fillna("").astype(str).str.upper().str.strip()
    desc = df["Description"].fillna("").astype(str)
    cp = df["Counterparty_Norm"].fillna("").astype(str).str.upper().str.strip()
    return typ.eq("CASH") | cp.eq("CASH") | desc.str.contains(r"\bCASH\b.*\b(?:WDL|WITHDRAWAL)\b", case=False, regex=True)

def _repeated_cluster_ids(cash_df: pd.DataFrame, count: int, days: int) -> set[str]:
    """Mark every transaction that belongs to a qualifying same-account cluster.

    This is intentionally cluster based rather than trailing-window based: if a
    group contains five withdrawals in 30 days, the early withdrawals in that
    same group are also part of the pattern.
    """
    marked: set[str] = set()
    if count <= 1:
        return set(cash_df["Txn_ID"])
    for _, g in cash_df.sort_values("Date").groupby("Account", dropna=False):
        g = g.sort_values("Date").reset_index(drop=True)
        for start in range(len(g)):
            # Every count-sized block that fits inside the time window qualifies.
            end = start + count - 1
            if end >= len(g):
                break
            if (g.loc[end, "Date"] - g.loc[start, "Date"]).days <= days:
                # Extend through any additional withdrawals still inside the same window.
                last = end
                while last + 1 < len(g) and (g.loc[last + 1, "Date"] - g.loc[start, "Date"]).days <= days:
                    last += 1
                marked.update(g.loc[start:last, "Txn_ID"].astype(str))
    return marked

def rule_R04_cash_withdrawals(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Cash withdrawal above threshold and/or part of a repeated-withdrawal cluster."""
    th = ctx.thresholds
    cash = df[_cash_mask(df) & (df["Money_Out"] > 0)].copy()
    if cash.empty:
        return []

    # The repeated-pattern check uses material withdrawals only. This prevents a
    # small routine cash withdrawal from being swept into a cluster simply
    # because it happened near several large withdrawals.
    material = cash[cash["Money_Out"] >= th.cash_withdrawal].copy()
    repeated = _repeated_cluster_ids(material, th.cash_repeat_count, th.cash_repeat_days)
    sym = currency_symbol(ctx.currency)

    hits = []
    for r in cash.itertuples():
        reasons = []
        if r.Money_Out >= th.cash_withdrawal:
            reasons.append(f"Cash withdrawal above {fmt_money(th.cash_withdrawal, sym)}")
        if str(r.Txn_ID) in repeated:
            reasons.append(f"Repeated cash withdrawals in {th.cash_repeat_days} days")
        if reasons:
            hits.append(Hit(r.Txn_ID, "R04", "; ".join(reasons)))
    return hits

def rule_R09_unidentified_transfer(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Outgoing transfer to a bank account not recognised as an own account."""
    typ = df["Type"].fillna("").astype(str).str.upper().str.strip()
    desc = df["Description"].fillna("").astype(str)
    transfer = typ.eq("TFR") | desc.str.match(r"^\s*TFR\b", case=False, na=False)
    m = transfer & (df["Money_Out"] > 0) & (~df["Is_Own_Account_Transfer"])

    hits = []
    for r in df[m].itertuples():
        sort_code = "" if pd.isna(r.Counterparty_Sort_Code) else str(r.Counterparty_Sort_Code).strip()
        account = "" if pd.isna(r.Counterparty_Account) else str(r.Counterparty_Account).strip()
        # CSV inference can turn an 8-digit account into "55820019.0".
        if re.fullmatch(r"\d+\.0", account):
            account = account[:-2]
        target = " ".join(x for x in (sort_code, account) if x)
        reason = f"Transfer to unidentified account {target}" if target else "Transfer to unidentified account"
        hits.append(Hit(r.Txn_ID, "R09", reason))
    return hits

def rule_R05_round_amount(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Large outgoing payment that is an exact multiple of the configured step."""
    th = ctx.thresholds
    amt = df["Money_Out"].astype(float)
    # Floating-safe divisibility test.
    nearest = (amt / th.round_step).round() * th.round_step
    is_round = (amt - nearest).abs() < 0.005
    m = (amt >= th.round_min) & is_round & (~df["Is_Own_Account_Transfer"])
    sym = currency_symbol(ctx.currency)
    return [Hit(r.Txn_ID, "R05", f"Round-amount payment ({fmt_money(r.Money_Out, sym)})")
            for r in df[m].itertuples()]

def rule_R06_new_counterparty(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """First observed high-value outgoing payment to a previously unseen counterparty.

    This is a screening heuristic, not a statement that the commercial
    relationship was legally "new". Known parties, own accounts and TFR account
    transfers are excluded because they have dedicated context/rules.
    """
    th = ctx.thresholds
    work = df[(df["Date"] >= ctx.analysis_start) & (df["Money_Out"] > 0)].copy()
    work["_cp"] = work["Counterparty_Norm"].map(_norm)
    work = work[work["_cp"] != ""]
    if work.empty:
        return []

    known = set(_known_party_maps(ctx).index)
    typ = work["Type"].fillna("").astype(str).str.upper().str.strip()
    work = work[(~work["Is_Own_Account_Transfer"]) & (~work["_cp"].isin(known)) & (~typ.eq("TFR"))]
    if work.empty:
        return []

    work = work.sort_values(["Date", "Source_File", "Source_Row"])
    first = work.groupby("_cp", as_index=False, sort=False).first()
    m = first["Money_Out"] >= th.new_cp_first_payment
    if th.new_cp_window_days > 0:
        window_start = ctx.insolvency_date - pd.Timedelta(days=th.new_cp_window_days)
        m &= first["Date"].between(window_start, ctx.insolvency_date, inclusive="both")
    first = first[m]

    sym = currency_symbol(ctx.currency)
    return [Hit(r.Txn_ID, "R06",
                f"New counterparty within analysis period; first payment {fmt_money(r.Money_Out, sym)}")
            for r in first.itertuples()]

def rule_R12_post_appointment(df: pd.DataFrame, ctx: Context) -> list[Hit]:
    """Transaction dated after appointment date - control marker, not an allegation."""
    m = df["Date"] > ctx.appointment_date
    return [Hit(r.Txn_ID, "R12", "Post-appointment transaction - verify authority and purpose") for r in df[m].itertuples()]

# scope: "analysis_period" - evaluated on transactions inside the analysis period (start..anchor);
#        "all"             - evaluated on all available data (connected parties and unidentified
#                            transfers matter whenever they occur, including after appointment);
#        "post_appointment" - looks after the appointment date regardless of the analysis period.
RULES = [
    (rule_R03_high_value, "analysis_period"),
    (rule_R12_post_appointment, "post_appointment"),
    (rule_R01_connected_party, "all"),
    (rule_R10_description_terms, "analysis_period"),
    (rule_R04_cash_withdrawals, "analysis_period"),
    (rule_R09_unidentified_transfer, "all"),
    (rule_R05_round_amount, "analysis_period"),
    (rule_R06_new_counterparty, "analysis_period"),
]

def time_band_series(dates: pd.Series, anchor: pd.Timestamp) -> pd.Series:
    """Calendar-month bands relative to the anchor (boundaries at anchor - 6/12/24/36 months)."""
    b6, b12, b24, b36 = (anchor - pd.DateOffset(months=m) for m in (6, 12, 24, 36))
    out = pd.Series("36m+", index=dates.index, dtype=object)
    out[dates >= b36] = "24-36m"
    out[dates >= b24] = "12-24m"
    out[dates >= b12] = "6-12m"
    out[dates >= b6] = "0-6m"
    out[dates > anchor] = "post-anchor"
    return out

# ----------------------------------------------------------------------------- engine
def run_rules(df: pd.DataFrame, ctx: Context, rules=RULES) -> pd.DataFrame:
    anchor = ctx.anchor_date if ctx.anchor_date is not None else ctx.appointment_date
    in_period = (df["Date"] >= ctx.analysis_start) & (df["Date"] <= anchor)
    hits: list[Hit] = []
    for item in rules:
        rule, scope = item if isinstance(item, tuple) else (item, "analysis_period")
        subset = df[in_period] if scope == "analysis_period" else df
        hits.extend(rule(subset, ctx))
    out = df.copy()
    out["In_Analysis_Period"] = in_period
    out["Days_To_Anchor"] = (anchor - out["Date"]).dt.days
    out["Time_Band"] = time_band_series(out["Date"], anchor)
    if hits:
        h = pd.DataFrame([vars(x) for x in hits]).drop_duplicates(["txn_id", "rule_id"])
        agg = h.groupby("txn_id").agg(Rule_IDs=("rule_id", lambda s: ";".join(sorted(s))),
                                      Review_Reasons=("reason", lambda s: "; ".join(s)))
        out = out.merge(agg, left_on="Txn_ID", right_index=True, how="left")
    else:
        out["Rule_IDs"] = None; out["Review_Reasons"] = None
    out["Review_Flag"] = out["Rule_IDs"].notna().map({True: "Y", False: "N"})
    return out

def counterparty_summary(out: pd.DataFrame, ctx: Context) -> pd.DataFrame:
    rel = ctx.known_parties.drop_duplicates("Name_Norm").set_index("Name_Norm")["Relationship"]
    g = out[~out["Is_Own_Account_Transfer"]].assign(_n=out["Counterparty_Norm"].map(_norm))
    s = g.groupby("Counterparty_Norm").agg(
        Paid_In=("Money_In", "sum"), Paid_Out=("Money_Out", "sum"), Txn_Count=("Txn_ID", "count"),
        First_Txn=("Date", "min"), Last_Txn=("Date", "max"),
        Flagged_Count=("Review_Flag", lambda x: int((x == "Y").sum())),
        Flagged_Amount=("Money_Out", lambda x: float(x[g.loc[x.index, "Review_Flag"] == "Y"].sum())),
    ).reset_index()
    s["Net"] = s["Paid_In"] - s["Paid_Out"]
    s["Relationship"] = s["Counterparty_Norm"].map(_norm).map(rel).fillna("")
    cols = ["Counterparty_Norm", "Relationship", "Paid_In", "Paid_Out", "Net", "Txn_Count", "First_Txn", "Last_Txn", "Flagged_Count", "Flagged_Amount"]
    return s[cols].sort_values("Paid_Out", ascending=False)

def case_summary(out: pd.DataFrame, ctx: Context, integrity: pd.DataFrame | None) -> pd.DataFrame:
    rows = [("Company", ctx.company), ("Accounts", out["Account"].nunique()), ("Banks", ", ".join(sorted(out["Bank"].unique()))),
            ("Period", f"{out['Date'].min():%d %b %Y} - {out['Date'].max():%d %b %Y}"),
            ("Procedure", ctx.procedure or "-"),
            ("Insolvency date", f"{ctx.insolvency_date:%d %b %Y}"), ("Appointment date", f"{ctx.appointment_date:%d %b %Y}"),
            ("Analysis anchor", f"{ctx.analysis_anchor}" + (f" ({ctx.anchor_date:%d %b %Y})" if ctx.anchor_date is not None else "")),
            ("Analysis lookback", (f"configured {ctx.lookback_months} months; start overridden" if ctx.start_overridden and ctx.lookback_months
                                   else ("custom start date" if ctx.start_overridden else f"{ctx.lookback_months} months"))),
            ("Analysis period", f"{ctx.analysis_start:%d %b %Y} - {(ctx.anchor_date or ctx.appointment_date):%d %b %Y} (analyst-defined; not a statutory look-back)"),
            ("Review output mode", ctx.review_output_mode),
            ("Transactions", len(out)), ("Total in", fmt_total(round(out["Money_In"].sum(), 2))), ("Total out", fmt_total(round(out["Money_Out"].sum(), 2))),
            ("Own-account transfers (excluded from counterparty view)", int(out["Is_Own_Account_Transfer"].sum())),
            ("Transactions flagged for review", int((out["Review_Flag"] == "Y").sum())),
            ("Flagged amount (money out)", fmt_total(round(out.loc[out["Review_Flag"] == "Y", "Money_Out"].sum(), 2)))]
    if out["Date"].min() > ctx.analysis_start:
        rows.append(("Data coverage note", f"Data coverage starts {out['Date'].min():%d %b %Y}, after analysis period start {ctx.analysis_start:%d %b %Y}"))
    if out["Rule_IDs"].notna().any():
        counts = out["Rule_IDs"].dropna().str.split(";").explode().value_counts().sort_index()
        rows += [(f"  {k}", int(v)) for k, v in counts.items()]
    if integrity is not None:
        ok = integrity["status"].eq("PASS").all()
        rows.append(("Data integrity checks", "PASS" if ok else "FAIL - see integrity sheet"))
    return pd.DataFrame(rows, columns=["Item", "Value"])

REVIEW_COLS = ["Date", "Bank", "Account", "Counterparty_Norm", "Relationship", "Money_Out", "Money_In",
               "Rule_IDs", "Review_Reasons", "Time_Band", "Days_To_Anchor", "In_Analysis_Period", "Description", "Type", "Balance", "Source_File", "Source_Row", "Txn_ID"]

def _style_workbook(book, sheet_names):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '£#,##0.00;[Red]-£#,##0.00;-'
    for name in sheet_names:
        ws = book[name]
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = border
        ws.row_dimensions[1].height = 30
        headers = [c.value for c in ws[1]]
        for col in ws.iter_cols(min_row=2):
            h = headers[col[0].column - 1]
            width = 12
            for cell in col:
                cell.font = Font(name="Arial", size=10); cell.border = border
                if h in ("Money_In", "Money_Out", "Amount_Signed", "Balance", "Paid_In", "Paid_Out", "Net", "Flagged_Amount", "Total in", "Total out", "Closing balance"):
                    cell.number_format = money
                if h in ("Date", "First_Txn", "Last_Txn"):
                    cell.number_format = "DD/MM/YYYY"
                if h in ("Review_Reasons", "Description"):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                width = max(width, min(70, len(str(cell.value)) + 2 if cell.value is not None else 0))
            ws.column_dimensions[col[0].column_letter].width = 46 if h in ("Review_Reasons",) else (40 if h == "Description" else min(width, 28))
        if "Rule_IDs" in headers:
            idx = headers.index("Rule_IDs") + 1
            letter = ws.cell(row=1, column=idx).column_letter

            # Review_Items uses restrained, practitioner-facing highlighting.
            # Only Rule_IDs and Review_Reasons are coloured; the rest of the
            # working paper stays white and easy to scan.
            if name == "Review_Items" and "Review_Reasons" in headers:
                reason_idx = headers.index("Review_Reasons") + 1
                reason_letter = ws.cell(row=1, column=reason_idx).column_letter
                neutral_fill = PatternFill("solid", fgColor="EEF3F8")  # soft blue-grey
                connected_fill = PatternFill("solid", fgColor="FCE8E6")  # soft rose
                post_fill = PatternFill("solid", fgColor="FFF2CC")  # soft yellow

                # Use direct cell fills rather than conditional formatting so the
                # practitioner sees the same restrained highlighting in every
                # Excel viewer and in exported screenshots.
                for row in range(2, ws.max_row + 1):
                    rule_ids = str(ws.cell(row=row, column=idx).value or "")
                    if "R12" in rule_ids:
                        fill = post_fill
                    elif "R01" in rule_ids:
                        fill = connected_fill
                    elif rule_ids:
                        fill = neutral_fill
                    else:
                        fill = PatternFill(fill_type=None)

                    ws.cell(row=row, column=idx).fill = fill
                    ws.cell(row=row, column=reason_idx).fill = fill
                    ws.cell(row=row, column=idx).font = Font(name="Arial", size=10, bold=True, color="1F3864")

                    reason = str(ws.cell(row=row, column=reason_idx).value or "")
                    description = str(ws.cell(row=row, column=headers.index("Description") + 1).value or "") if "Description" in headers else ""
                    longest = max(len(reason), len(description))
                    ws.row_dimensions[row].height = 24 if longest <= 75 else (36 if longest <= 145 else 48)
            else:
                rng = f"A2:{ws.cell(row=1, column=len(headers)).column_letter}{ws.max_row}"
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("R01",${letter}2))'], fill=PatternFill("solid", fgColor="F8CBAD")))
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("R12",${letter}2))'], fill=PatternFill("solid", fgColor="FFE699")))
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEN(${letter}2)>0'], fill=PatternFill("solid", fgColor="FFF2CC")))
        ws.auto_filter.ref = ws.dimensions

def _integrity_summary(integ: pd.DataFrame) -> pd.DataFrame:
    """Return a practitioner-facing reconciliation summary.

    The full expected/actual comparison remains available on the hidden
    ``Integrity_Detail`` worksheet.  This sheet intentionally shows only the
    values a practitioner needs to confirm that each source statement was
    reconciled successfully.
    """
    return pd.DataFrame({
        "Bank": integ["bank"],
        "Rows": integ["rows_actual"].astype(int),
        "Total in": integ["total_in_actual"],
        "Total out": integ["total_out_actual"],
        "Closing balance": integ["closing_balance_actual"],
        "Status": integ["status"],
    })


def write_output(out: pd.DataFrame, ctx: Context, path: Path, integrity_path: Path | None):
    integ = pd.read_csv(integrity_path) if integrity_path and Path(integrity_path).exists() else None
    rel = ctx.known_parties.drop_duplicates("Name_Norm").set_index("Name_Norm")["Relationship"]
    out = out.copy()
    out["Relationship"] = out["Counterparty_Norm"].map(_norm).map(rel).fillna("")
    review = out[out["Review_Flag"] == "Y"]
    if ctx.review_output_mode.lower().startswith("analysis"):
        review = review[review["In_Analysis_Period"] | review["Rule_IDs"].fillna("").str.contains("R12|R01|R09")]
    review = review.sort_values(["Money_Out", "Date"], ascending=[False, True])
    review = review[[c for c in REVIEW_COLS if c in review.columns]]
    tx_cols = [c for c in REVIEW_COLS if c in out.columns] + [c for c in out.columns if c not in REVIEW_COLS]
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
        case_summary(out, ctx, integ).to_excel(w, index=False, sheet_name="Case_Summary")
        review.to_excel(w, index=False, sheet_name="Review_Items")
        counterparty_summary(out, ctx).to_excel(w, index=False, sheet_name="Counterparty_Summary")
        out[tx_cols].to_excel(w, index=False, sheet_name="Transactions")
        if integ is not None:
            _integrity_summary(integ).to_excel(w, index=False, sheet_name="Integrity")
            integ.to_excel(w, index=False, sheet_name="Integrity_Detail")
        _style_workbook(w.book, [n for n in ["Case_Summary", "Review_Items", "Counterparty_Summary", "Transactions", "Integrity", "Integrity_Detail"] if n in w.book.sheetnames])
        cs = w.book["Case_Summary"]; cs.column_dimensions["A"].width = 52; cs.column_dimensions["B"].width = 36
        if "Integrity" in w.book.sheetnames:
            iw = w.book["Integrity"]
            for col, width in {"A": 14, "B": 12, "C": 18, "D": 18, "E": 20, "F": 12}.items():
                iw.column_dimensions[col].width = width
        if "Integrity_Detail" in w.book.sheetnames:
            w.book["Integrity_Detail"].sheet_state = "hidden"

def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--transactions", required=True); ap.add_argument("--config", required=True)
    ap.add_argument("--integrity", default=None); ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    ctx = load_context(Path(a.config))
    df = load_transactions(Path(a.transactions))
    out = run_rules(df, ctx)
    write_output(out, ctx, Path(a.out), Path(a.integrity) if a.integrity else None)
    n = int((out["Review_Flag"] == "Y").sum())
    print(f"{ctx.company}: {len(out)} transactions, {n} flagged for review -> {a.out}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
