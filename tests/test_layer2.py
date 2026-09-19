"""Layer 2 tests: rule validation against expected test cases (synthetic ground truth)."""
import sys
from pathlib import Path
import pandas as pd
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "review"))
import review_engine as re_

TX = ROOT / "output" / "normalized_transactions.csv"
CFG = ROOT / "synthetic" / "data" / "case_config.xlsx"
EXP = ROOT / "synthetic" / "data" / "expected_flags.csv"

def _run():
    ctx = re_.load_context(CFG)
    df = re_.load_transactions(TX)
    return re_.run_rules(df, ctx), ctx

def _match(out, exp_row):
    """Locate an expected planted transaction by bank / date / amount (money out)."""
    m = (out["Bank"] == exp_row.account) & (out["Date"] == pd.Timestamp(exp_row.date)) & \
        ((out["Money_Out"] - abs(exp_row.amount)).abs() < 0.005)
    return out[m]

def _check_rule(rule_id):
    out, _ = _run()
    exp = pd.read_csv(EXP)
    exp = exp[exp["rule"] == rule_id]
    assert len(exp) > 0, f"no expected flags for {rule_id}"
    missing = []
    for r in exp.itertuples():
        hit = _match(out, r)
        assert len(hit) >= 1, f"planted transaction not found for {r.flag_id}"
        if not hit["Rule_IDs"].fillna("").str.contains(rule_id).any():
            missing.append(r.flag_id)
    assert not missing, f"{rule_id} missed expected items: {missing}"
    return out

def test_engine_accumulates_multiple_rules():
    out, _ = _run()
    multi = out[out["Rule_IDs"].fillna("").str.contains(";")]
    assert len(multi) > 0
    for v in out["Rule_IDs"].dropna():
        assert all(p.startswith("R") for p in v.split(";"))

def test_R12_post_appointment_exactly_expected():
    out = _check_rule("R12")
    ctx = re_.load_context(CFG)
    post = out[out["Date"] > ctx.appointment_date]
    assert len(post) == 2, f"expected exactly 2 post-appointment transactions, got {len(post)}"
    assert (post["Rule_IDs"].fillna("").str.contains("R12")).all()

def test_R03_high_value_finds_expected():
    _check_rule("R03")

def test_R03_excludes_routine_payees():
    out, ctx = _run()
    hits = out[out["Rule_IDs"].fillna("").str.contains("R03")]
    assert not hits["Counterparty_Norm"].str.contains("HMRC", case=False).any(), "HMRC payments must not be flagged as high-value"
    assert not hits["Is_Own_Account_Transfer"].any(), "own-account transfers must not be flagged"

def test_R01_connected_party_finds_expected_and_does_not_guess_initials():
    out = _check_rule("R01")
    # J SMITH is deliberately ambiguous because both John Smith and Jane Smith
    # exist in Known_Parties. R01 must not infer identity from an initial.
    ambiguous = out[out["Counterparty_Norm"].eq("J SMITH")]
    assert len(ambiguous) > 0
    assert not ambiguous["Rule_IDs"].fillna("").str.contains("R01").any()

def test_F21_accumulates_R01_and_R12():
    out, _ = _run()
    row = out[(out["Bank"] == "Lloyds") & (out["Date"] == pd.Timestamp("2026-03-18")) &
              ((out["Money_Out"] - 4000).abs() < 0.005)]
    assert len(row) == 1
    ids = set(row.iloc[0]["Rule_IDs"].split(";"))
    assert {"R01", "R12"}.issubset(ids)

def test_R10_description_terms_finds_expected():
    out = _check_rule("R10")
    hits = out[out["Rule_IDs"].fillna("").str.contains("R10")]
    assert len(hits) == 2

def test_R04_cash_withdrawals_finds_expected_cluster_only():
    out = _check_rule("R04")
    hits = out[out["Rule_IDs"].fillna("").str.contains("R04")]
    assert len(hits) == 6
    assert hits["Review_Reasons"].str.contains("Repeated cash withdrawals in 30 days").all()
    assert (hits["Money_Out"] >= 1000).all()

def test_R09_unidentified_transfer_finds_expected():
    out = _check_rule("R09")
    hits = out[out["Rule_IDs"].fillna("").str.contains("R09")]
    assert len(hits) == 1
    assert "40-11-07 55820019" in hits.iloc[0]["Review_Reasons"]

def test_R05_round_amount_finds_expected():
    _check_rule("R05")

def test_R06_new_counterparty_finds_expected_and_uses_recency_window():
    out = _check_rule("R06")
    hits = out[out["Rule_IDs"].fillna("").str.contains("R06")]
    # Vantage (F16) and Delta (F18) are the two ordinary counterparties whose
    # first material payment falls inside the configured 180-day window.
    assert set(hits["Counterparty_Norm"]) == {"VANTAGE PROPERTY SERVICES LTD", "DELTA ASSET PARTNERS LLP"}


def test_flagged_set_equals_planted_set():
    """README claim: 21 of 21 planted items flagged; no other transactions flagged."""
    out, _ = _run()
    exp = pd.read_csv(EXP)
    planted = set()
    for r in exp.itertuples():
        hit = _match(out, r)
        assert len(hit) == 1, f"planted item {r.flag_id} matched {len(hit)} rows"
        planted.add(hit["Txn_ID"].iloc[0])
    flagged = set(out.loc[out["Review_Flag"] == "Y", "Txn_ID"])
    assert flagged == planted, f"extra: {flagged - planted}; missing: {planted - flagged}"


def test_analysis_period_is_computed_from_anchor_and_lookback():
    ctx = re_.load_context(CFG)
    assert ctx.analysis_anchor == "Appointment"
    assert ctx.lookback_months == 24
    assert ctx.analysis_start == ctx.appointment_date - pd.DateOffset(months=24)


def test_time_bands_and_scope():
    out, ctx = _run()
    assert set(out["Time_Band"].unique()) <= {"0-6m", "6-12m", "12-24m", "24-36m", "36m+", "post-anchor"}
    post = out[out["Date"] > ctx.appointment_date]
    assert (post["Time_Band"] == "post-anchor").all()
    # rules scoped to the analysis period must not fire outside it (R03/R04/R05/R06/R10)
    outside = out[~out["In_Analysis_Period"] & (out["Date"] <= ctx.appointment_date)]
    assert not outside["Rule_IDs"].fillna("").str.contains("R03|R04|R05|R06|R10").any()


# ---- period model on non-default configurations (in-memory Case series, no files needed)
def _case(**kw):
    base = {"Company name": "X", "Appointment date": pd.Timestamp("2026-03-14"), "Insolvency date": pd.Timestamp("2026-03-10"),
            "Petition_Date": None, "Resolution_Date": pd.Timestamp("2026-03-10"), "Analysis_Anchor": "Appointment",
            "Analysis_Anchor_Date": None, "Analysis_Lookback_Months": 24, "Analysis_Period_Start_Override": None}
    base.update(kw)
    return pd.Series(base)

def test_period_ends_exactly_on_anchor_for_resolution_and_petition():
    _, ctx = _run()
    df = re_.load_transactions(TX)
    for anchor_name, anchor_dt in [("Resolution", pd.Timestamp("2026-03-10")), ("Petition", pd.Timestamp("2026-02-01"))]:
        ctx2 = re_.Context(**{**vars(ctx), "analysis_anchor": anchor_name, "anchor_date": anchor_dt,
                               "analysis_start": anchor_dt - pd.DateOffset(months=24)})
        out = re_.run_rules(df, ctx2)
        after = out[out["Date"] > anchor_dt]
        assert not after["In_Analysis_Period"].any(), f"{anchor_name}: transactions after anchor inside period"
        assert (after["Time_Band"] == "post-anchor").all()

def test_custom_anchor_uses_analysis_anchor_date():
    name, anchor, months, start, ov = re_.resolve_analysis_period(_case(Analysis_Anchor="Custom", Analysis_Anchor_Date=pd.Timestamp("2025-12-31")))
    assert name == "Custom" and anchor == pd.Timestamp("2025-12-31") and start == pd.Timestamp("2023-12-31")

def test_override_changes_start_but_not_anchor():
    name, anchor, months, start, ov = re_.resolve_analysis_period(
        _case(Analysis_Anchor="Resolution", Analysis_Lookback_Months=None, Analysis_Period_Start_Override=pd.Timestamp("2025-01-01")))
    assert name == "Resolution" and anchor == pd.Timestamp("2026-03-10") and start == pd.Timestamp("2025-01-01") and ov

def test_period_validation_rejects_bad_inputs():
    import pytest
    with pytest.raises(ValueError):
        re_.resolve_analysis_period(_case(Analysis_Lookback_Months=-12))
    with pytest.raises(ValueError):
        re_.resolve_analysis_period(_case(Analysis_Period_Start_Override=pd.Timestamp("2027-01-01")))
    with pytest.raises(ValueError):
        re_.resolve_analysis_period(_case(Analysis_Anchor="Custom"))  # no Analysis_Anchor_Date

def test_time_band_uses_calendar_month_boundaries():
    anchor = pd.Timestamp("2026-08-31")
    s = pd.Series(pd.to_datetime(["2026-02-28", "2026-02-27", "2025-08-31", "2025-08-30", "2026-09-01"]))
    assert list(re_.time_band_series(s, anchor)) == ["0-6m", "6-12m", "6-12m", "12-24m", "post-anchor"]

def test_review_output_mode_analysis_period_only():
    out, ctx = _run()
    ctx2 = re_.Context(**{**vars(ctx), "review_output_mode": "Analysis period only"})
    import tempfile, os
    p = Path(tempfile.mkdtemp()) / "ro.xlsx"
    re_.write_output(out, ctx2, p, None)
    r = pd.read_excel(p, sheet_name="Review_Items")
    assert (r["In_Analysis_Period"] | r["Rule_IDs"].str.contains("R12|R01|R09")).all()

def test_integrity_workbook_has_compact_visible_and_hidden_detail():
    """Final workbook keeps a concise practitioner view and a hidden technical reconciliation."""
    out, ctx = _run()
    import tempfile
    from openpyxl import load_workbook
    p = Path(tempfile.mkdtemp()) / "integrity_view.xlsx"
    integrity_path = ROOT / "output" / "normalized_transactions_integrity.csv"
    re_.write_output(out, ctx, p, integrity_path)

    visible = pd.read_excel(p, sheet_name="Integrity")
    assert list(visible.columns) == ["Bank", "Rows", "Total in", "Total out", "Closing balance", "Status"]
    assert len(visible) == 3
    assert visible["Status"].eq("PASS").all()

    wb = load_workbook(p, read_only=False, data_only=True)
    assert "Integrity_Detail" in wb.sheetnames
    assert wb["Integrity_Detail"].sheet_state == "hidden"
    detail_headers = [c.value for c in wb["Integrity_Detail"][1]]
    assert {"rows_expected", "rows_actual", "total_in_expected", "total_in_actual", "status"}.issubset(detail_headers)
